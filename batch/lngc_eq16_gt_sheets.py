# Colour the two lngc_eq16 ground-truth sheets by what the pipeline caught.
# Every row and column of the source files is kept unchanged and in its original order; the source files are only read.
# Four columns are appended: caught, what_was_caught, caught_on_sheets, how_read.  Green row = caught.
# A row is caught when its instrument tag (P&ID tag from FLOC, else the PI / external tag) was extracted on the 10 sheets:
# same letters + number (A/B suffix ignored), or letters that differ only by trailing letters (list LIC, drawn LICA).
import csv, os, re
from collections import defaultdict
import openpyxl
from openpyxl.styles import PatternFill, Font

SRC = os.path.expanduser("~/Downloads/pids_bhcp/")
OUT = "lngc_eq16_output/eval/lngc_eq16_ground_truth_vs_extracted.xlsx"
GREEN = PatternFill("solid", fgColor="C6EFCE")
strip = lambda n: re.sub(r"[A-Z]$", "", n)

ext = defaultdict(list)
for m in csv.DictReader(open("lngc_eq16_output/mastersheet.csv")):
    t = re.sub(r"^1?-?111-|^1-", "", m["tag"]).replace("TIC-A-", "TICA-")
    x = re.match(r"([A-Z]{1,6})-?(\d{4})([A-Z]?)", t)
    if x: ext[(x.group(1), x.group(2) + x.group(3))].append(m)
short = lambda s: re.sub(r"_\d_\d_?\d? ?|_p0$", " ", s[19:]).strip()

FLOC = re.compile(r"^111([A-Z]{1,6})-(\d{4})([A-Z]?)(?:-|$)")
DCS = re.compile(r"CA:KIT:111([A-Z]+?)(\d{4})([A-Z]?)(?![0-9])")
def caught(key):
    if not key: return []
    L, n = key
    hits = [m for k, ms in ext.items() if k[0] == L and strip(k[1]) == strip(n) for m in ms]
    if not hits:   # list and drawing letters differ only at the end (LIC/LICA, PDIA/PDIAH, UZ/UZV)
        hits = [m for k, ms in ext.items() if strip(k[1]) == strip(n) and (k[0].startswith(L) or L.startswith(k[0])) for m in ms]
    return hits
def extra(hits):
    return ["YES" if hits else "NO", " | ".join(sorted({h["tag"] for h in hits})),
            " | ".join(sorted({short(h["sheet"]) for h in hits})), " | ".join(sorted({h["source"] for h in hits}))]
EXTRA = ["caught", "what_was_caught", "caught_on_sheets", "how_read"]

wb = openpyxl.Workbook(); wb.remove(wb.active)

# ---- PI_TAGS_LNGC_eq16.csv: all rows
pi_rows = list(csv.reader(open(SRC + "PI_TAGS_LNGC_eq16.csv", encoding="utf-8-sig")))
hdr = pi_rows[0]; i_name, i_floc = hdr.index("PI Tag Name"), hdr.index("FLOC")
dcs2floc = {}
ws = wb.create_sheet("PI_TAGS_LNGC_eq16"); ws.append(hdr + EXTRA); n_pi = 0
for row in pi_rows[1:]:
    f, d = FLOC.match(row[i_floc]), DCS.match(row[i_name])
    key = (f.group(1), f.group(2) + f.group(3)) if f else (d.group(1), d.group(2) + d.group(3)) if d else None
    if f and d: dcs2floc[(d.group(1), d.group(2) + d.group(3))] = key
    hits = caught(key); n_pi += bool(hits)
    ws.append(row + extra(hits))
    if hits:
        for c in ws[ws.max_row]: c.fill = GREEN

# ---- Flash GC sensors for Anurav.xlsx: all sheets, all rows
src = openpyxl.load_workbook(SRC + "Flash GC sensors for Anurav.xlsx", data_only=True); n_fgc = 0
for sws in src.worksheets:
    rows = list(sws.iter_rows(values_only=True))
    hdr = list(rows[0]); i_ext = hdr.index("external_tag")
    ws = wb.create_sheet(f"Flash GC sensors - {sws.title}"[:31]); ws.append(hdr + EXTRA)
    for row in rows[1:]:
        d = DCS.match(str(row[i_ext] or ""))
        key = (d.group(1), d.group(2) + d.group(3)) if d else None
        hits = caught(dcs2floc.get(key, key)); n_fgc += bool(hits)
        ws.append(list(row) + extra(hits))
        if hits:
            for c in ws[ws.max_row]: c.fill = GREEN

for ws in wb.worksheets:
    for c in ws[1]: c.font = Font(bold=True)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
wb.save(OUT)
print(f"PI_TAGS: {len(pi_rows) - 1} rows, {n_pi} green | Flash GC: {sum(s.max_row - 1 for s in src.worksheets)} rows, {n_fgc} green -> {OUT}")
