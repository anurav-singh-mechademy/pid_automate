# Batch tag extraction over a folder of P&IDs (loose PDFs + PDFs inside .zip archives), resumable and duplicate-safe.
#   python batch/batch_tags.py pids_bhcp pids_bhcp_output
# 1. unzip every archive into <out>/_sources/<zip name>/
# 2. inventory every PDF page; a page is processed once, keyed by its rendered content (72 dpi grayscale MD5), so the
#    same sheet under another file name / in another zip / inside a multi-page PDF is skipped as a duplicate.
#    Same file name with different content (another revision) is processed under a suffixed output name.
# 3. per page: tags_<name>_p<n>.csv + _overlay.png in <out>/sheets/, one ledger row in <out>/ledger.csv
#    (re-running skips every page already in the ledger with status ok)
# 4. mastersheet: all tags of all processed pages -> <out>/mastersheet.csv and .xlsx (all rows + unique tags per sheet)
import csv, hashlib, re, sys, time, traceback, zipfile
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from tag_pipeline import *            # noqa: F401,F403 - the notebook cell's functions and constants
import tag_pipeline as TP

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
SHEETS, SOURCES = OUT / "sheets", OUT / "_sources"
LEDGER, INVENTORY = OUT / "ledger.csv", OUT / "inventory.csv"
LEDGER_COLS = ["out_name", "file", "page", "fingerprint", "status", "seconds", "rotation", "bubble_mm", "size_method",
               "bubbles", "tags", "tags_grammar_ok", "error"]
for d in (SHEETS, SOURCES):
    d.mkdir(parents=True, exist_ok=True)

# ---------- 1. unzip ----------
for z in sorted(SRC.glob("*.zip")):
    dest = SOURCES / z.stem
    if not dest.exists():
        with zipfile.ZipFile(z) as zf:
            for m in zf.infolist():
                if m.filename.lower().endswith(".pdf") and "__MACOSX" not in m.filename:
                    p = dest / Path(m.filename).name
                    p.parent.mkdir(parents=True, exist_ok=True)
                    p.write_bytes(zf.read(m))

# ---------- 2. inventory + duplicate detection ----------
is_pdf = lambda p: p.is_file() and p.suffix.lower() == ".pdf"           # .pdf and .PDF
pdfs = sorted(filter(is_pdf, SRC.iterdir())) + sorted(filter(is_pdf, SOURCES.rglob("*")))
inventory, seen_fp, name_fps = [], {}, defaultdict(list)
for pdf in pdfs:
    try:
        doc = pymupdf.open(pdf)
    except Exception as e:
        inventory.append(dict(file=str(pdf), page="", fingerprint="", status="unreadable", duplicate_of="", out_name="", note=str(e)))
        continue
    for i, pg in enumerate(doc):
        pix = pg.get_pixmap(dpi=72, colorspace=pymupdf.csGRAY, alpha=False)
        fp = hashlib.md5(pix.samples).hexdigest()
        base = re.sub(r"[^\w\-. &]+", "_", pdf.stem).strip()
        if fp in seen_fp:
            inventory.append(dict(file=str(pdf), page=i, fingerprint=fp, status="duplicate", duplicate_of=seen_fp[fp]["out_name"],
                                  out_name="", note=f"same page as {seen_fp[fp]['file']} p{seen_fp[fp]['page']}"))
            continue
        key = (base.lower(), i)
        note = ""
        out_name = f"{base}_p{i}"
        if name_fps[key]:                         # same name + page, different drawing content (another revision)
            out_name = f"{base}__{pdf.parent.name}_p{i}"
            note = f"same file name as {name_fps[key][0]['file']} but different content"
        row = dict(file=str(pdf), page=i, fingerprint=fp, status="unique", duplicate_of="", out_name=out_name, note=note)
        name_fps[key].append(row); seen_fp[fp] = row; inventory.append(row)
with open(INVENTORY, "w", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=list(inventory[0])); w.writeheader(); w.writerows(inventory)
todo = [r for r in inventory if r["status"] == "unique"]
print(f"{len(pdfs)} PDFs, {len(inventory)} pages: {len(todo)} unique, "
      f"{sum(r['status'] == 'duplicate' for r in inventory)} duplicates, {sum(r['status'] == 'unreadable' for r in inventory)} unreadable",
      flush=True)

# ---------- 3. run the pipeline per unique page ----------
done = set()
if LEDGER.exists():
    done = {r["fingerprint"] for r in csv.DictReader(open(LEDGER)) if r["status"] == "ok"}
else:
    with open(LEDGER, "w", newline="") as fh:
        csv.DictWriter(fh, fieldnames=LEDGER_COLS).writeheader()
ocr_engine = RapidOCR()
rf_model, rf_names = load_rfdetr(WEIGHTS)
for n, r in enumerate(todo, 1):
    if r["fingerprint"] in done:
        continue
    t0 = time.time()
    led = dict(out_name=r["out_name"], file=r["file"], page=r["page"], fingerprint=r["fingerprint"])
    try:
        page = load_page(Path(r["file"]), r["page"])
        rotation, _ = detect_orientation(page) if AUTO_ORIENT else (0, {})
        if rotation:
            page = page.rotate(rotation, expand=True)
        fragments = ocr_page(page, ocr_engine)
        bubble_px, size_method, _ = measure_bubble(page, rf_model, rf_names, fragments)
        bubbles = detect_multiscale(page, rf_model, rf_names, choose_tile(bubble_px))
        attach(fragments, bubbles)
        tags = build_tags(fragments, bubbles)
        out_csv = SHEETS / f"tags_{r['out_name']}.csv"
        with open(out_csv, "w", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(tags[0]) if tags else ["tag"]); w.writeheader(); w.writerows(tags)
        draw_overlay(page, bubbles, tags).save(SHEETS / f"tags_{r['out_name']}_overlay.png")
        led.update(status="ok", rotation=rotation, bubble_mm=round(bubble_px / MM, 1) if bubble_px else "", size_method=size_method,
                   bubbles=len(bubbles), tags=len(tags), tags_grammar_ok=sum(t["grammar_ok"] for t in tags), error="")
    except Exception as e:
        led.update(status="error", error=f"{type(e).__name__}: {e}")
        traceback.print_exc()
    led["seconds"] = round(time.time() - t0)
    with open(LEDGER, "a", newline="") as fh:
        csv.DictWriter(fh, fieldnames=LEDGER_COLS).writerow(led)
    print(f"[{n}/{len(todo)}] {r['out_name']}: {led['status']} {led.get('tags', '')} tags, {led['seconds']}s", flush=True)

# ---------- 4. mastersheet ----------
import openpyxl
from openpyxl.styles import Font
ledger = {r["out_name"]: r for r in csv.DictReader(open(LEDGER)) if r["status"] == "ok"}
rows = []
for name, L in ledger.items():
    f = SHEETS / f"tags_{name}.csv"
    if f.exists():
        for t in csv.DictReader(open(f)):
            if t.get("tag"):
                rows.append(dict(sheet=name, source_file=Path(L["file"]).name, source_path=L["file"], page=L["page"], **t))
cols = ["sheet", "source_file", "page", "tag", "source", "confidence", "grammar_ok", "lines", "ocr_conf", "det_conf",
        "bubble_label", "dist_mm", "bbox", "source_path"]
with open(OUT / "mastersheet.csv", "w", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore"); w.writeheader(); w.writerows(rows)
uniq = defaultdict(lambda: dict(sheets=set(), n=0, conf=0.0, sources=set(), ok=False))
for t in rows:
    u = uniq[t["tag"]]; u["sheets"].add(t["sheet"]); u["n"] += 1; u["sources"].add(t["source"])
    u["conf"] = max(u["conf"], float(t["confidence"])); u["ok"] |= t["grammar_ok"] == "True"
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "all tags"
ws.append(cols)
for t in rows:
    ws.append([t.get(c) for c in cols])
ws2 = wb.create_sheet("unique tags")
ws2.append(["tag", "grammar_ok", "max_confidence", "occurrences", "n_sheets", "sources", "sheets"])
for tag, u in sorted(uniq.items(), key=lambda kv: (not kv[1]["ok"], kv[0])):
    ws2.append([tag, u["ok"], round(u["conf"], 3), u["n"], len(u["sheets"]), " | ".join(sorted(u["sources"])), " | ".join(sorted(u["sheets"]))])
ws3 = wb.create_sheet("sheets")
ws3.append(LEDGER_COLS)
for r in csv.DictReader(open(LEDGER)):
    ws3.append([r[c] for c in LEDGER_COLS])
for s in wb.worksheets:
    for c in s[1]:
        c.font = Font(bold=True)
    s.freeze_panes = "A2"
wb.save(OUT / "mastersheet.xlsx")
import tag_filter                                                  # precision filter: keep / drop_reason / kind per row
print("precision filter:", dict(tag_filter.filter_mastersheet(OUT)), flush=True)
print(f"mastersheet: {len(rows)} tag rows from {len(ledger)} pages, {len(uniq)} unique tags "
      f"({sum(u['ok'] for u in uniq.values())} pass grammar) -> {OUT}/mastersheet.xlsx", flush=True)
