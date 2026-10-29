# pids_bhcp batch mastersheet vs the plant-wide Cove Point sensor export (BHCP ground truth).
#   python batch/eval_pids_bhcp.py
# Keys: list internal_tag "5TI-5090A_CAL" / extracted "TI-5090A" -> (letters, loop number + suffix).
#   exact      = same letters + number + suffix
#   no suffix  = same letters + number, A/B suffix ignored (lists often write PI-5122 for drawn PI-5122A/B)
#   same loop  = same first letter + number (alarm/indicator letters differ: TI vs TIT vs TAH)
# Agreement = share of extracted tags found in the list. Recall needs a denominator of list instruments that are drawn
# on these sheets; the list covers the whole plant, so recall is reported (a) over the whole list (lower bound) and
# (b) over the 57 sheets with a PDF text layer, where "printed on the sheet" comes from the PDF text, independent of OCR.
import csv, re, sys
from collections import defaultdict, Counter
from pathlib import Path
import pymupdf, openpyxl
from openpyxl.styles import PatternFill, Font

OUT = Path("pids_bhcp_output"); EV = OUT / "eval"; EV.mkdir(exist_ok=True)
LIST = "bhcp/dominion/Cove_Point_LNG_Sensors_export_2026-10-15_15-22-57.csv"
ALLOW_DIGIT_SUFFIX = True     # count A1/B2-suffixed tags that the tag cell still marks as format fails

P = re.compile(r"^(?:[A-Z]+:)?(\d{1,2})?([A-Z]{1,5})-?(\d{3,5})([A-Z]\d?)?(?![0-9])")
plant = defaultdict(list)
for r in csv.DictReader(open(LIST, encoding="utf-8-sig")):
    m = P.match(r["internal_tag"])
    if m and len(m.group(2)) > 1:            # single letters (K-511A, E-512B) are equipment keys, not instruments
        plant[(m.group(2), m.group(3) + (m.group(4) or ""))].append(r)
strip = lambda n: re.sub(r"[A-Z]\d?$", "", n)
plant_nosuf = defaultdict(set); plant_loop = defaultdict(set)
for k in plant:
    plant_nosuf[(k[0], strip(k[1]))].add(k); plant_loop[(k[0][0], strip(k[1]))].add(k)

E = re.compile(r"(?:\d{1,3}-)?([A-Z]{2,5})-(\d{3,5}(?:[A-Z]\d?)?)(?:-\d+)?$")
ext = defaultdict(list)
for r in csv.DictReader(open(OUT / "mastersheet.csv")):
    m = E.match(r["tag"])
    ok = r["grammar_ok"] == "True" or (ALLOW_DIGIT_SUFFIX and m and re.search(r"\d[A-Z]\d$", m.group(2)))
    if m and ok:
        ext[(m.group(1), m.group(2))].append(r)

ext_nosuf = defaultdict(list)
for e in ext:
    ext_nosuf[(e[0], strip(e[1]))].append(e)
def match(k):
    if k in plant: return "EXACT"
    if (k[0], strip(k[1])) in plant_nosuf: return "NO-SUFFIX"
    if (k[0][0], strip(k[1])) in plant_loop: return "SAME LOOP"
    return "NOT IN LIST"
def category(L):
    if L.startswith("HV") and len(L) > 2: return "hand valve on an instrument (HVPSV, HVPIT...)"
    if L in ("HV", "VV", "CK", "TP"): return "local device"
    if L.endswith("Y"): return "solenoid / relay / converter"
    if L.startswith("ZS") or L[-2:] in ("SH", "SL") or L.endswith(("SHH", "SLL")): return "switch"
    if L.endswith(("T", "E")): return "transmitter / element"
    if L.endswith("V"): return "valve"
    if L.endswith(("I", "A", "C", "IC")): return "indicator / alarm / controller"
    return "other"

status = {k: match(k) for k in ext}
cnt = Counter(status.values()); n = len(ext)

# ---- recall (a): whole list
def found(k):     # list key -> extracted keys matching it at the loosest level that still keeps the letters
    if k in ext: return "EXACT", [k]
    hits = ext_nosuf.get((k[0], strip(k[1])), [])
    return ("NO-SUFFIX", hits) if hits else ("MISSED", [])
gt_status = {k: found(k) for k in plant}

# ---- recall (b): text-layer sheets, list keys printed in the PDF text
inv = [r for r in csv.DictReader(open(OUT / "inventory.csv")) if r["status"] == "unique"]
printed = defaultdict(set)
for r in inv:
    t = pymupdf.open(r["file"])[int(r["page"])].get_text()
    if len(t) < 200: continue
    flat = re.sub(r"\s+", " ", t.upper())
    for mm in re.finditer(r"(?<![A-Z0-9])(\d{0,2})([A-Z]{2,5})[- ]?(\d{3,5})([A-Z]\d?)?(?![0-9])", flat):
        k = (mm.group(2), mm.group(3) + (mm.group(4) or ""))
        if k in plant: printed[k].add(r["out_name"])
text_sheets = {r["out_name"] for r in inv if len(pymupdf.open(r["file"])[int(r["page"])].get_text()) >= 200}
ext_on = lambda k, sheets: any(x["sheet"] in sheets for e in ext_nosuf.get((k[0], strip(k[1])), []) for x in ext[e])
scope_hit = {k: ext_on(k, printed[k]) for k in printed}

# ---- tables
ex_rows = []
for k in sorted(ext):
    lst = sorted({r["internal_tag"] for kk in (plant_loop.get((k[0][0], strip(k[1]))) or []) for r in plant[kk]})
    for r in ext[k]:
        ex_rows.append(dict(status=status[k], key=f"{k[0]}-{k[1]}", extracted_tag=r["tag"], sheet=r["sheet"], source_file=r["source_file"],
                            source=r["source"], confidence=r["confidence"], category="" if status[k] != "NOT IN LIST" else category(k[0]),
                            list_tags=" | ".join(lst)[:250], list_name=plant[next(iter(plant_loop[(k[0][0], strip(k[1]))]))][0]["name"] if lst else "",
                            bbox=r["bbox"]))
gt_rows = []
for k in sorted(plant):
    st, hits = gt_status[k]
    gt_rows.append(dict(status=st, key=f"{k[0]}-{k[1]}", list_tag=plant[k][0]["internal_tag"], list_name=plant[k][0]["name"],
                        list_rows=len(plant[k]), extracted=" | ".join(sorted({x["tag"] for h in hits for x in ext[h]}))[:200],
                        sheets=" | ".join(sorted({x["sheet"] for h in hits for x in ext[h]}))[:300],
                        printed_in_pdf_text=" | ".join(sorted(printed.get(k, [])))))
for name, rows in (("list_to_extracted", gt_rows), ("extracted_to_list", ex_rows)):
    with open(EV / f"{name}.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
FILL = {"EXACT": "C6EFCE", "NO-SUFFIX": "E2F0D9", "SAME LOOP": "FFEB9C", "MISSED": "FFC7CE", "NOT IN LIST": "FFC7CE"}
wb = openpyxl.Workbook(); wb.remove(wb.active)
for title, rows in (("extracted -> list", ex_rows), ("list -> extracted", gt_rows)):
    ws = wb.create_sheet(title); cols = list(rows[0]); ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in cols])
        for c in ws[ws.max_row]: c.fill = PatternFill("solid", fgColor=FILL.get(r["status"], "FFFFFF"))
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
wb.save(EV / "pids_bhcp_vs_list.xlsx")

# ---- summary
print(f"list: {len(plant)} instrument keys | extracted: {n} distinct keys ({sum(len(v) for v in ext.values())} rows) on 260 pages")
print("AGREEMENT (extracted keys found in the list):")
acc = 0
for s in ("EXACT", "NO-SUFFIX", "SAME LOOP"):
    acc += cnt[s]; print(f"  {s:10s} {cnt[s]:5d}   cumulative {acc}/{n} = {acc / n:.0%}")
print(f"  NOT IN LIST {cnt['NOT IN LIST']} by kind:", dict(Counter(category(k[0]) for k in ext if status[k] == "NOT IN LIST").most_common()))
g = Counter(s for s, _ in gt_status.values())
print(f"RECALL over the whole plant list (lower bound): exact {g['EXACT']}, +no-suffix {g['EXACT'] + g['NO-SUFFIX']}/{len(plant)} "
      f"= {(g['EXACT'] + g['NO-SUFFIX']) / len(plant):.0%}")
print(f"RECALL on the {len(text_sheets)} text-layer sheets (list keys printed in PDF text): "
      f"{sum(scope_hit.values())}/{len(scope_hit)} = {sum(scope_hit.values()) / max(1, len(scope_hit)):.0%}")
print("  missed there:", sorted(f"{k[0]}-{k[1]}" for k, v in scope_hit.items() if not v)[:60])
lp = Counter(k[0] for k, (s, _) in gt_status.items() if s == "MISSED")
print("list keys never extracted, by letters:", lp.most_common(15))
print("files:", sorted(p.name for p in EV.iterdir()))
