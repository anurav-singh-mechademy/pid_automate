# Every extracted tag of lngc_eq16 with its precision verdict. Red = false positive.
#   in the ground-truth files               -> correct (green)
#   not in the files, real drawn instrument -> correct (light green), checked by eye
#   interlock label / reference / not a tag / misread -> FALSE POSITIVE (red), checked by eye
import csv, re
import openpyxl
from openpyxl.styles import PatternFill, Font
E = "lngc_eq16_output/eval/"
verdict = {r["key"]: r["verdict"] for r in csv.DictReader(open(E + "not_in_PI_TAGS_verdicts.csv"))}
CAT = {"real instrument": ("correct: real instrument on the sheet, not in the GT files", "E2F0D9"),
       "interlock": ("FALSE POSITIVE: interlock (UZ) label, not an instrument", "FFC7CE"),
       "reference": ("FALSE POSITIVE: reference in a note / off-sheet connector", "FFC7CE"),
       "not a tag": ("FALSE POSITIVE: not a tag", "FFC7CE"),
       "misread": ("FALSE POSITIVE: misread", "FFC7CE")}
rows = []
for r in csv.DictReader(open(E + "extracted_to_PI_TAGS.csv")):
    if r["status"] != "NOT IN PI_TAGS":
        v, fill = "correct: in the ground-truth files" + (" (letters differ)" if r["status"] == "same loop" else ""), "C6EFCE"
    else:
        v, fill = next(CAT[k] for k in CAT if verdict[r["key"]].startswith(k))
        if "misread" in v: v += " (" + verdict[r["key"]].split("(")[1]
    rows.append(dict(verdict=v, extracted_tag=r["extracted_tag"], key=r["key"], sheet=re.sub(r"_\d_\d_?\d? ?|_p0$", " ", r["sheet"][19:]).strip(),
                     how_read=r["source"], confidence=r["confidence"], bbox_px=r["bbox"], _fill=fill))
order = lambda r: (not r["verdict"].startswith("FALSE"), r["verdict"], r["key"])
rows.sort(key=order)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "extracted tags"
cols = [c for c in rows[0] if c != "_fill"]; ws.append(cols)
for r in rows:
    ws.append([r[c] for c in cols])
    for c in ws[ws.max_row]: c.fill = PatternFill("solid", fgColor=r["_fill"])
keys = {}
for r in rows: keys.setdefault(r["key"], r["verdict"])
s = wb.create_sheet("summary", 0); s.append(["verdict", "distinct tags", "rows"])
for v in sorted({r["verdict"] for r in rows}, key=lambda v: (not v.startswith("FALSE"), v)):
    s.append([v, sum(x == v for x in keys.values()), sum(r["verdict"] == v for r in rows)])
fp = sum(v.startswith("FALSE") for v in keys.values())
s.append([]); s.append(["precision (distinct tags)", f"{len(keys) - fp}/{len(keys)} = {(len(keys) - fp) / len(keys):.0%}"])
for w in wb.worksheets:
    for c in w[1]: c.font = Font(bold=True)
    w.freeze_panes = "A2"
    for i, c in enumerate(w[1], 1): w.column_dimensions[c.column_letter].width = 60 if i == 1 else 22
ws.auto_filter.ref = ws.dimensions
wb.save(E + "lngc_eq16_precision_false_positives.xlsx")
print(f"{len(rows)} rows, {len(keys)} distinct, {fp} false positives")
